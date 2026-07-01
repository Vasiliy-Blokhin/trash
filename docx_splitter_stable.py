from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from copy import deepcopy
import re
import os
import shutil
import tempfile
from typing import List, Dict, Any, Optional, Tuple, Union, Callable
from dataclasses import dataclass, field


@dataclass
class RunStyle:
    text: str = ""
    bold: Optional[bool] = None
    italic: Optional[bool] = None
    underline: Optional[Union[bool, int]] = None
    font_name: Optional[str] = None
    font_size: Optional[float] = None
    font_color: Optional[Tuple[int, int, int]] = None
    highlight_color: Optional[int] = None
    all_caps: Optional[bool] = None
    small_caps: Optional[bool] = None
    strike: Optional[bool] = None
    double_strike: Optional[bool] = None
    subscript: Optional[bool] = None
    superscript: Optional[bool] = None
    shadow: Optional[bool] = None
    outline: Optional[bool] = None
    emboss: Optional[bool] = None
    imprint: Optional[bool] = None
    character_style: Optional[str] = None


@dataclass
class ParagraphStyle:
    style_name: Optional[str] = None
    alignment: Optional[int] = None
    space_before: Optional[float] = None
    space_after: Optional[float] = None
    line_spacing: Optional[float] = None
    line_spacing_rule: Optional[int] = None
    left_indent: Optional[float] = None
    right_indent: Optional[float] = None
    first_line_indent: Optional[float] = None
    keep_together: Optional[bool] = None
    keep_with_next: Optional[bool] = None
    page_break_before: Optional[bool] = None
    widow_control: Optional[bool] = None
    outline_level: Optional[int] = None


@dataclass
class LogicalBlock:
    text: str
    word_count: int
    paragraph_style: ParagraphStyle
    runs: List[RunStyle]
    block_type: str
    body_element_indices: List[int] = field(default_factory=list)

    def __repr__(self):
        preview = self.text[:80] + "..." if len(self.text) > 80 else self.text
        return f"LogicalBlock(type={self.block_type}, words={self.word_count}, text='{preview}')"


class DocxBlockSplitter:
    def __init__(self, min_words: int = 10, merge_same_style: bool = True,
                 preserve_headings: bool = True, preserve_list_items: bool = False):
        self.min_words = min_words
        self.merge_same_style = merge_same_style
        self.preserve_headings = preserve_headings
        self.preserve_list_items = preserve_list_items

    def _count_words(self, text: str) -> int:
        if not text or not text.strip():
            return 0
        words = re.findall(r'[\w\u0400-\u04FF]+', text)
        return len(words)

    def _extract_run_style(self, run) -> RunStyle:
        rs = RunStyle(text=run.text)
        rs.bold = run.bold
        rs.italic = run.italic
        rs.underline = run.underline
        rs.font_name = run.font.name
        rs.font_size = run.font.size.pt if run.font.size else None
        if run.font.color and run.font.color.rgb:
            rs.font_color = (run.font.color.rgb[0], run.font.color.rgb[1], run.font.color.rgb[2])
        rs.all_caps = run.font.all_caps
        rs.small_caps = run.font.small_caps
        rs.strike = run.font.strike
        rs.double_strike = run.font.double_strike
        rs.subscript = run.font.subscript
        rs.superscript = run.font.superscript
        rs.shadow = run.font.shadow
        rs.outline = run.font.outline
        rs.emboss = run.font.emboss
        rs.imprint = run.font.imprint
        if run.style and run.style.name:
            rs.character_style = run.style.name
        try:
            rPr = run._r.find(qn('w:rPr'))
            if rPr is not None:
                highlight = rPr.find(qn('w:highlight'))
                if highlight is not None:
                    val = highlight.get(qn('w:val'))
                    if val:
                        from docx.enum.text import WD_COLOR_INDEX
                        rs.highlight_color = getattr(WD_COLOR_INDEX, val.upper(), None)
        except:
            pass
        return rs

    def _extract_paragraph_style(self, paragraph) -> ParagraphStyle:
        ps = ParagraphStyle()
        if paragraph.style and paragraph.style.name:
            ps.style_name = paragraph.style.name
        pf = paragraph.paragraph_format
        ps.alignment = pf.alignment
        ps.space_before = pf.space_before.pt if pf.space_before else None
        ps.space_after = pf.space_after.pt if pf.space_after else None
        ps.line_spacing = pf.line_spacing
        ps.line_spacing_rule = pf.line_spacing_rule
        ps.left_indent = pf.left_indent.pt if pf.left_indent else None
        ps.right_indent = pf.right_indent.pt if pf.right_indent else None
        ps.first_line_indent = pf.first_line_indent.pt if pf.first_line_indent else None
        ps.keep_together = pf.keep_together
        ps.keep_with_next = pf.keep_with_next
        ps.page_break_before = pf.page_break_before
        ps.widow_control = pf.widow_control
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    val = outline.get(qn('w:val'))
                    if val is not None:
                        ps.outline_level = int(val)
        except:
            pass
        return ps

    def _styles_compatible(self, style1: ParagraphStyle, style2: ParagraphStyle) -> bool:
        if style1.style_name != style2.style_name:
            return False
        if style1.alignment != style2.alignment:
            return False

        def approx_equal(a, b, tolerance=0.5):
            if a is None and b is None:
                return True
            if a is None or b is None:
                return False
            return abs(a - b) < tolerance

        if not approx_equal(style1.left_indent, style2.left_indent):
            return False
        if not approx_equal(style1.right_indent, style2.right_indent):
            return False
        if not approx_equal(style1.first_line_indent, style2.first_line_indent):
            return False
        if not approx_equal(style1.space_before, style2.space_before):
            return False
        if not approx_equal(style1.space_after, style2.space_after):
            return False
        return True

    def _is_heading(self, paragraph) -> bool:
        if paragraph.style and paragraph.style.name:
            style_name = paragraph.style.name.lower()
            if any(x in style_name for x in ['heading', 'заголовок', 'title', 'head']):
                return True
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    return True
        except:
            pass
        return False

    def _is_list_item(self, paragraph) -> bool:
        if paragraph.style and paragraph.style.name:
            style_name = paragraph.style.name.lower()
            if any(x in style_name for x in ['list', 'bullet', 'number', 'список']):
                return True
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    return True
        except:
            pass
        return False

    def _extract_paragraph_style_from_xml(self, p_elem) -> ParagraphStyle:
        ps = ParagraphStyle()
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val:
                        ps.style_name = val
                jc = pPr.find(qn('w:jc'))
                if jc is not None:
                    val = jc.get(qn('w:val'))
                    if val:
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        ps.alignment = getattr(WD_ALIGN_PARAGRAPH, val.upper(), None)
                spacing = pPr.find(qn('w:spacing'))
                if spacing is not None:
                    before = spacing.get(qn('w:before'))
                    after = spacing.get(qn('w:after'))
                    line = spacing.get(qn('w:line'))
                    line_rule = spacing.get(qn('w:lineRule'))
                    if before is not None:
                        ps.space_before = int(before) / 20.0
                    if after is not None:
                        ps.space_after = int(after) / 20.0
                    if line is not None:
                        ps.line_spacing = int(line) / 20.0
                    if line_rule is not None:
                        from docx.enum.text import WD_LINE_SPACING
                        ps.line_spacing_rule = getattr(WD_LINE_SPACING, line_rule.upper(), None)
                ind = pPr.find(qn('w:ind'))
                if ind is not None:
                    left = ind.get(qn('w:left'))
                    right = ind.get(qn('w:right'))
                    first_line = ind.get(qn('w:firstLine'))
                    if left is not None:
                        ps.left_indent = int(left) / 20.0
                    if right is not None:
                        ps.right_indent = int(right) / 20.0
                    if first_line is not None:
                        ps.first_line_indent = int(first_line) / 20.0
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    val = outline.get(qn('w:val'))
                    if val is not None:
                        ps.outline_level = int(val)
                for prop_name, attr_name in [
                    ('w:keepLines', 'keep_together'),
                    ('w:keepNext', 'keep_with_next'),
                    ('w:pageBreakBefore', 'page_break_before'),
                    ('w:widowControl', 'widow_control'),
                ]:
                    prop = pPr.find(qn(prop_name))
                    if prop is not None:
                        val = prop.get(qn('w:val'))
                        if val is not None:
                            setattr(ps, attr_name, val == '1' or val.lower() == 'true')
        except:
            pass
        return ps

    def _is_heading_from_xml(self, p_elem) -> bool:
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    return True
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val and any(x in val.lower() for x in ['heading', 'заголовок', 'title', 'head']):
                        return True
        except:
            pass
        return False

    def _is_list_item_from_xml(self, p_elem) -> bool:
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    return True
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val and any(x in val.lower() for x in ['list', 'bullet', 'number', 'список']):
                        return True
        except:
            pass
        return False

    def _apply_paragraph_style(self, new_para, style: ParagraphStyle):
        if style.style_name:
            try:
                new_para.style = style.style_name
            except:
                pass
        pf = new_para.paragraph_format
        if style.alignment is not None:
            pf.alignment = style.alignment
        if style.space_before is not None:
            pf.space_before = Pt(style.space_before)
        if style.space_after is not None:
            pf.space_after = Pt(style.space_after)
        if style.line_spacing is not None:
            pf.line_spacing = style.line_spacing
        if style.line_spacing_rule is not None:
            pf.line_spacing_rule = style.line_spacing_rule
        if style.left_indent is not None:
            pf.left_indent = Pt(style.left_indent)
        if style.right_indent is not None:
            pf.right_indent = Pt(style.right_indent)
        if style.first_line_indent is not None:
            pf.first_line_indent = Pt(style.first_line_indent)
        if style.keep_together is not None:
            pf.keep_together = style.keep_together
        if style.keep_with_next is not None:
            pf.keep_with_next = style.keep_with_next
        if style.page_break_before is not None:
            pf.page_break_before = style.page_break_before
        if style.widow_control is not None:
            pf.widow_control = style.widow_control

    def split(self, docx_path: str) -> List[LogicalBlock]:
        doc = Document(docx_path)
        blocks = []
        raw_blocks = []

        body = doc.element.body
        body_elements = list(body)
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

        for idx, element in enumerate(body_elements):
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag

            if tag == 'tbl':
                texts = []
                for t in element.findall('.//w:t', namespaces=nsmap):
                    if t.text:
                        texts.append(t.text)
                text = ''.join(texts).strip()
                word_count = self._count_words(text)

                raw_blocks.append({
                    'text': text,
                    'word_count': word_count,
                    'para_style': ParagraphStyle(),
                    'runs': [],
                    'type': 'table',
                    'element_index': idx,
                    'is_table': True,
                })

            elif tag == 'p':
                para = None
                for p in doc.paragraphs:
                    if p._p is element:
                        para = p
                        break

                if para is None:
                    continue

                text = para.text.strip()
                word_count = self._count_words(text)
                para_style = self._extract_paragraph_style(para)
                runs = [self._extract_run_style(run) for run in para.runs]

                if self._is_heading(para):
                    block_type = 'heading'
                elif self._is_list_item(para):
                    block_type = 'list_item'
                else:
                    block_type = 'paragraph'

                raw_blocks.append({
                    'text': text,
                    'word_count': word_count,
                    'para_style': para_style,
                    'runs': runs,
                    'type': block_type,
                    'element_index': idx,
                    'is_table': False,
                })

        i = 0
        while i < len(raw_blocks):
            current = raw_blocks[i]

            if current.get('is_table'):
                blocks.append(LogicalBlock(
                    text=current['text'],
                    word_count=current['word_count'],
                    paragraph_style=current['para_style'],
                    runs=current['runs'],
                    block_type='table',
                    body_element_indices=[current['element_index']]
                ))
                i += 1
                continue

            if self.preserve_headings and current['type'] == 'heading' and current['word_count'] > 0:
                blocks.append(LogicalBlock(
                    text=current['text'],
                    word_count=current['word_count'],
                    paragraph_style=current['para_style'],
                    runs=current['runs'],
                    block_type='heading',
                    body_element_indices=[current['element_index']]
                ))
                i += 1
                continue

            if self.preserve_list_items and current['type'] == 'list_item' and current['word_count'] > 0:
                blocks.append(LogicalBlock(
                    text=current['text'],
                    word_count=current['word_count'],
                    paragraph_style=current['para_style'],
                    runs=current['runs'],
                    block_type='list_item',
                    body_element_indices=[current['element_index']]
                ))
                i += 1
                continue

            merged_text = current['text']
            merged_words = current['word_count']
            merged_para_style = current['para_style']
            merged_runs = list(current['runs'])
            merged_type = current['type']
            merged_indices = [current['element_index']]

            j = i + 1
            while merged_words < self.min_words and j < len(raw_blocks):
                next_block = raw_blocks[j]

                if next_block.get('is_table'):
                    break
                if next_block['type'] == 'heading' and self.preserve_headings:
                    break
                if next_block['type'] == 'list_item' and self.preserve_list_items:
                    break
                if self.merge_same_style and not self._styles_compatible(merged_para_style, next_block['para_style']):
                    break

                if merged_text:
                    merged_text += '\n'
                merged_text += next_block['text']
                merged_words += next_block['word_count']
                merged_runs.extend(next_block['runs'])
                merged_indices.append(next_block['element_index'])
                if merged_type != next_block['type']:
                    merged_type = 'merged'
                j += 1

            blocks.append(LogicalBlock(
                text=merged_text,
                word_count=merged_words,
                paragraph_style=merged_para_style,
                runs=merged_runs,
                block_type=merged_type,
                body_element_indices=merged_indices
            ))
            i = j if j > i + 1 else i + 1

        return blocks

    def _copy_docx_with_rels(self, source_path: str, dest_path: str) -> None:
        shutil.copy2(source_path, dest_path)

    def create_block_document(self, block: LogicalBlock, output_path: str,
                               source_docx_path: str = None) -> None:
        if source_docx_path and os.path.exists(source_docx_path):
            self._copy_docx_with_rels(source_docx_path, output_path)
            doc = Document(output_path)

            body = doc.element.body
            keep_indices = set(block.body_element_indices)

            all_elements = list(body)
            for idx in range(len(all_elements) - 1, -1, -1):
                if idx not in keep_indices:
                    body.remove(all_elements[idx])

            doc.save(output_path)
        else:
            new_doc = Document()
            for idx in block.body_element_indices:
                if idx < len(body_elements := list(new_doc.element.body)):
                    new_doc.element.body.append(deepcopy(body_elements[idx]))
            new_doc.save(output_path)

    def split_to_files(self, docx_path: str, output_prefix: str) -> List[str]:
        blocks = self.split(docx_path)
        output_paths = []
        for idx, block in enumerate(blocks):
            output_path = f"{output_prefix}_block_{idx:03d}.docx"
            self.create_block_document(block, output_path, source_docx_path=docx_path)
            output_paths.append(output_path)
        return output_paths

    def _process_paragraph_element(self, para, text_processor, skip_headings, skip_list_items, nsmap):
        text = ''.join(t.text or '' for t in para.findall('.//w:t', namespaces=nsmap))
        para_style = self._extract_paragraph_style_from_xml(para)
        is_heading = self._is_heading_from_xml(para)
        is_list_item = self._is_list_item_from_xml(para)

        should_process = True
        if skip_headings and is_heading:
            should_process = False
        if skip_list_items and is_list_item:
            should_process = False

        if should_process:
            para_info = {
                'is_heading': is_heading,
                'is_list_item': is_list_item,
                'style_name': para_style.style_name,
                'alignment': para_style.alignment,
            }
            new_text = text_processor(text, para_info)
        else:
            new_text = text

        t_elements = para.findall('.//w:t', namespaces=nsmap)
        if t_elements:
            original_lengths = []
            for t in t_elements:
                original_lengths.append(len(t.text) if t.text else 0)
            for t in t_elements:
                t.text = ''
            self._replace_text_in_t_elements(t_elements, new_text, original_lengths)

    def process_document(self, input_path: str, output_path: str,
                         text_processor: Callable[[str, Dict[str, Any]], str],
                         skip_headings: bool = True,
                         skip_list_items: bool = False) -> None:
        self._copy_docx_with_rels(input_path, output_path)
        doc = Document(output_path)
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        all_paragraphs = doc.element.body.findall('.//w:p', namespaces=nsmap)
        for para in all_paragraphs:
            self._process_paragraph_element(para, text_processor, skip_headings, skip_list_items, nsmap)
        doc.save(output_path)

    def _replace_text_in_t_elements(self, t_elements: List, new_text: str, original_lengths: List[int]) -> None:
        if not t_elements:
            return
        new_text_len = len(new_text)
        if new_text_len == 0:
            return
        num_elements = len(t_elements)
        total_original = sum(original_lengths)
        if total_original == 0 or num_elements == 1:
            t_elements[0].text = new_text
            if new_text.endswith(' ') or new_text.startswith(' '):
                t_elements[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            return
        portions = []
        for length in original_lengths:
            portion = round(new_text_len * length / total_original)
            portions.append(portion)
        diff = sum(portions) - new_text_len
        if diff > 0:
            for _ in range(diff):
                max_idx = portions.index(max(portions))
                portions[max_idx] -= 1
        elif diff < 0:
            for _ in range(-diff):
                max_idx = portions.index(max(portions))
                portions[max_idx] += 1
        current_pos = 0
        for i, t in enumerate(t_elements):
            portion_length = portions[i]
            portion_text = new_text[current_pos:current_pos + portion_length]
            current_pos += portion_length
            t.text = portion_text
            if portion_text and (portion_text.endswith(' ') or portion_text.startswith(' ')):
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        assert current_pos == new_text_len, f"current_pos={current_pos} != new_text_len={new_text_len}"