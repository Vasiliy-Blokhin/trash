from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
import re
import os
import shutil
from typing import List, Dict, Any, Optional, Tuple, Union, Callable
from dataclasses import dataclass, field
import asyncio

# from Site.utils.text.random import text_random


@dataclass
class RunStyle:
    text: str = ""
    bold: Optional[bool] = None
    italic: Optional[bool] = None
    underline: Optional[Union[bool, int]] = None
    font_name: Optional[str] = None
    font_size: Optional[float] = None
    font_color: Optional[Tuple[int, int, int]] = None
    highlight_color: Optional[int] = None
    all_caps: Optional[bool] = None
    small_caps: Optional[bool] = None
    strike: Optional[bool] = None
    double_strike: Optional[bool] = None
    subscript: Optional[bool] = None
    superscript: Optional[bool] = None
    shadow: Optional[bool] = None
    outline: Optional[bool] = None
    emboss: Optional[bool] = None
    imprint: Optional[bool] = None
    character_style: Optional[str] = None


@dataclass
class ParagraphStyle:
    style_name: Optional[str] = None
    alignment: Optional[int] = None
    space_before: Optional[float] = None
    space_after: Optional[float] = None
    line_spacing: Optional[float] = None
    line_spacing_rule: Optional[int] = None
    left_indent: Optional[float] = None
    right_indent: Optional[float] = None
    first_line_indent: Optional[float] = None
    keep_together: Optional[bool] = None
    keep_with_next: Optional[bool] = None
    page_break_before: Optional[bool] = None
    widow_control: Optional[bool] = None
    outline_level: Optional[int] = None


@dataclass
class LogicalBlock:
    text: str
    word_count: int
    paragraph_style: ParagraphStyle
    runs: List[RunStyle]
    block_type: str
    body_element_indices: List[int] = field(default_factory=list)

    def __repr__(self):
        preview = self.text[:80] + "..." if len(self.text) > 80 else self.text
        return f"LogicalBlock(type={self.block_type}, words={self.word_count}, text='{preview}')"


# ---------------------------------------------------------------------------
# Markdown helpers
# ---------------------------------------------------------------------------

_MD_INLINE_RE = re.compile(
    r'(\*\*\*.+?\*\*\*|___.+?___'                                   # bold+italic
    r'|\*\*.+?\*\*|__.+?__'                                          # bold
    r'|~~.+?~~'                                                      # strikethrough
    r'|==.+?=='                                                       # highlight
    r'|\+\+.+?\+\+'                                                  # underline
    r'|~[^~]+?~'                                                      # subscript
    r'|\^[^\^]+?\^'                                                   # superscript
    r'|\*[^*]+?\*|_[^_]+?_'                                          # italic
    r'|`[^`]+?`'                                                      # inline code
    r'|\[.+?\]\(.+?\))',                                             # hyperlink [text](url)
    re.DOTALL,
)

_MD_HEADING_RE = re.compile(r'^(#{1,6})\s+(.*)$')
_MD_BLOCKQUOTE_RE = re.compile(r'^>\s?(.*)$')
_MD_HR_RE = re.compile(r'^(-{3,}|\*{3,}|_{3,})\s*$')
_MD_BULLET_RE = re.compile(r'^(\s*)([-*])\s+(.*)$')
_MD_NUMBERED_RE = re.compile(r'^(\s*)(\d+)\.\s+(.*)$')

_MD_TABLE_SEP_RE = re.compile(r'^\|?\s*[-:]+[-|\s:]*$')

_MD_CODE_FONT = 'Consolas'
_HYPERLINK_REL_TYPE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
_W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

_RPR_ORDER = [
    'w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps',
    'w:smallCaps', 'w:strike', 'w:dstrike', 'w:outline', 'w:shadow',
    'w:emboss', 'w:imprint', 'w:noProof', 'w:snapToGrid', 'w:vanish',
    'w:webHidden', 'w:color', 'w:spacing', 'w:w', 'w:kern', 'w:position',
    'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect', 'w:bdr', 'w:shd',
    'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang',
    'w:eastAsianLayout', 'w:specVanish', 'w:oMath',
]

_PPR_ORDER = [
    'w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore',
    'w:framePr', 'w:widowControl', 'w:numPr', 'w:suppressLineNumbers',
    'w:pBdr', 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku',
    'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE',
    'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
    'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
    'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
    'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
    'w:rPr', 'w:sectPr', 'w:pPrChange',
]


class DocxWorker:
    def __init__(self, min_words: int = 10, merge_same_style: bool = True,
                 preserve_headings: bool = True, preserve_list_items: bool = False,
                 base_font_name: str = 'Times New Roman',
                 base_font_size: float = 14.0,
                 base_alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
        self.min_words = min_words
        self.base_font_name = base_font_name
        self.base_font_size = float(base_font_size)
        self.base_alignment = base_alignment
        self.merge_same_style = merge_same_style
        self.preserve_headings = preserve_headings
        self.preserve_list_items = preserve_list_items

    def _count_words(self, text: str) -> int:
        if not text or not text.strip():
            return 0
        words = re.findall(r'[\w\u0400-\u04FF]+', text)
        return len(words)

    def _extract_run_style(self, run) -> RunStyle:
        rs = RunStyle(text=run.text)
        rs.bold = run.bold
        rs.italic = run.italic
        rs.underline = run.underline
        rs.font_name = run.font.name
        rs.font_size = run.font.size.pt if run.font.size else None
        if run.font.color and run.font.color.rgb:
            rs.font_color = (run.font.color.rgb[0], run.font.color.rgb[1], run.font.color.rgb[2])
        rs.all_caps = run.font.all_caps
        rs.small_caps = run.font.small_caps
        rs.strike = run.font.strike
        rs.double_strike = run.font.double_strike
        rs.subscript = run.font.subscript
        rs.superscript = run.font.superscript
        rs.shadow = run.font.shadow
        rs.outline = run.font.outline
        rs.emboss = run.font.emboss
        rs.imprint = run.font.imprint
        if run.style and run.style.name:
            rs.character_style = run.style.name
        try:
            rPr = run._r.find(qn('w:rPr'))
            if rPr is not None:
                highlight = rPr.find(qn('w:highlight'))
                if highlight is not None:
                    val = highlight.get(qn('w:val'))
                    if val:
                        from docx.enum.text import WD_COLOR_INDEX
                        rs.highlight_color = getattr(WD_COLOR_INDEX, val.upper(), None)
        except:
            pass
        return rs

    def _extract_paragraph_style(self, paragraph) -> ParagraphStyle:
        ps = ParagraphStyle()
        if paragraph.style and paragraph.style.name:
            ps.style_name = paragraph.style.name
        pf = paragraph.paragraph_format
        ps.alignment = pf.alignment
        ps.space_before = pf.space_before.pt if pf.space_before else None
        ps.space_after = pf.space_after.pt if pf.space_after else None
        ps.line_spacing = pf.line_spacing
        ps.line_spacing_rule = pf.line_spacing_rule
        ps.left_indent = pf.left_indent.pt if pf.left_indent else None
        ps.right_indent = pf.right_indent.pt if pf.right_indent else None
        ps.first_line_indent = pf.first_line_indent.pt if pf.first_line_indent else None
        ps.keep_together = pf.keep_together
        ps.keep_with_next = pf.keep_with_next
        ps.page_break_before = pf.page_break_before
        ps.widow_control = pf.widow_control
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    val = outline.get(qn('w:val'))
                    if val is not None:
                        ps.outline_level = int(val)
        except:
            pass
        return ps

    def _styles_compatible(self, style1: ParagraphStyle, style2: ParagraphStyle) -> bool:
        if style1.style_name != style2.style_name:
            return False
        if style1.alignment != style2.alignment:
            return False
        def approx_equal(a, b, tolerance=0.5):
            if a is None and b is None:
                return True
            if a is None or b is None:
                return False
            return abs(a - b) < tolerance
        if not approx_equal(style1.left_indent, style2.left_indent):
            return False
        if not approx_equal(style1.right_indent, style2.right_indent):
            return False
        if not approx_equal(style1.first_line_indent, style2.first_line_indent):
            return False
        if not approx_equal(style1.space_before, style2.space_before):
            return False
        if not approx_equal(style1.space_after, style2.space_after):
            return False
        return True

    def _is_heading(self, paragraph) -> bool:
        if paragraph.style and paragraph.style.name:
            style_name = paragraph.style.name.lower()
            if any(x in style_name for x in ['heading', 'title', 'head']):
                return True
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    return True
        except:
            pass
        return False

    def _is_list_item(self, paragraph) -> bool:
        if paragraph.style and paragraph.style.name:
            style_name = paragraph.style.name.lower()
            if any(x in style_name for x in ['list', 'bullet', 'number']):
                return True
        try:
            pPr = paragraph._p.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    return True
        except:
            pass
        return False

    def _extract_paragraph_style_from_xml(self, p_elem) -> ParagraphStyle:
        ps = ParagraphStyle()
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val:
                        ps.style_name = val
                jc = pPr.find(qn('w:jc'))
                if jc is not None:
                    val = jc.get(qn('w:val'))
                    if val:
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        ps.alignment = getattr(WD_ALIGN_PARAGRAPH, val.upper(), None)
                spacing = pPr.find(qn('w:spacing'))
                if spacing is not None:
                    before = spacing.get(qn('w:before'))
                    after = spacing.get(qn('w:after'))
                    line = spacing.get(qn('w:line'))
                    line_rule = spacing.get(qn('w:lineRule'))
                    if before is not None:
                        ps.space_before = int(before) / 20.0
                    if after is not None:
                        ps.space_after = int(after) / 20.0
                    if line is not None:
                        ps.line_spacing = int(line) / 20.0
                    if line_rule is not None:
                        from docx.enum.text import WD_LINE_SPACING
                        ps.line_spacing_rule = getattr(WD_LINE_SPACING, line_rule.upper(), None)
                ind = pPr.find(qn('w:ind'))
                if ind is not None:
                    left = ind.get(qn('w:left'))
                    right = ind.get(qn('w:right'))
                    first_line = ind.get(qn('w:firstLine'))
                    if left is not None:
                        ps.left_indent = int(left) / 20.0
                    if right is not None:
                        ps.right_indent = int(right) / 20.0
                    if first_line is not None:
                        ps.first_line_indent = int(first_line) / 20.0
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    val = outline.get(qn('w:val'))
                    if val is not None:
                        ps.outline_level = int(val)
                for prop_name, attr_name in [
                    ('w:keepLines', 'keep_together'),
                    ('w:keepNext', 'keep_with_next'),
                    ('w:pageBreakBefore', 'page_break_before'),
                    ('w:widowControl', 'widow_control'),
                ]:
                    prop = pPr.find(qn(prop_name))
                    if prop is not None:
                        val = prop.get(qn('w:val'))
                        if val is not None:
                            setattr(ps, attr_name, val == '1' or val.lower() == 'true')
        except:
            pass
        return ps

    def _is_heading_from_xml(self, p_elem) -> bool:
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                outline = pPr.find(qn('w:outlineLvl'))
                if outline is not None:
                    return True
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val and any(x in val.lower() for x in ['heading', 'title', 'head']):
                        return True
        except:
            pass
        return False

    def _is_list_item_from_xml(self, p_elem) -> bool:
        try:
            pPr = p_elem.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    return True
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    val = pStyle.get(qn('w:val'))
                    if val and any(x in val.lower() for x in ['list', 'bullet', 'number']):
                        return True
        except:
            pass
        return False

    def _apply_paragraph_style(self, new_para, style: ParagraphStyle):
        if style.style_name:
            try:
                new_para.style = style.style_name
            except:
                pass
        pf = new_para.paragraph_format
        if style.alignment is not None:
            pf.alignment = style.alignment
        if style.space_before is not None:
            pf.space_before = Pt(style.space_before)
        if style.space_after is not None:
            pf.space_after = Pt(style.space_after)
        if style.line_spacing is not None:
            pf.line_spacing = style.line_spacing
        if style.line_spacing_rule is not None:
            pf.line_spacing_rule = style.line_spacing_rule
        if style.left_indent is not None:
            pf.left_indent = Pt(style.left_indent)
        if style.right_indent is not None:
            pf.right_indent = Pt(style.right_indent)
        if style.first_line_indent is not None:
            pf.first_line_indent = Pt(style.first_line_indent)
        if style.keep_together is not None:
            pf.keep_together = style.keep_together
        if style.keep_with_next is not None:
            pf.keep_with_next = style.keep_with_next
        if style.page_break_before is not None:
            pf.page_break_before = style.page_break_before
        if style.widow_control is not None:
            pf.widow_control = style.widow_control

    # ------------------------------------------------------------------
    # Block building
    # ------------------------------------------------------------------

    def _build_raw_blocks(self, doc, nsmap) -> List[Dict[str, Any]]:
        raw_blocks = []
        body = doc.element.body
        for idx, element in enumerate(list(body)):
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag
            if tag == 'tbl':
                texts = []
                for t in element.findall('.//w:t', namespaces=nsmap):
                    if t.text:
                        texts.append(t.text)
                text = ''.join(texts).strip()
                raw_blocks.append({
                    'text': text,
                    'word_count': self._count_words(text),
                    'para_style': ParagraphStyle(),
                    'runs': [],
                    'type': 'table',
                    'element_index': idx,
                    'element': element,
                    'is_table': True,
                })
            elif tag == 'p':
                text = ''.join(t.text or '' for t in element.findall('.//w:t', namespaces=nsmap)).strip()
                para_style = self._extract_paragraph_style_from_xml(element)
                if self._is_heading_from_xml(element):
                    block_type = 'heading'
                elif self._is_list_item_from_xml(element):
                    block_type = 'list_item'
                else:
                    block_type = 'paragraph'
                raw_blocks.append({
                    'text': text,
                    'word_count': self._count_words(text),
                    'para_style': para_style,
                    'runs': [],
                    'type': block_type,
                    'element_index': idx,
                    'element': element,
                    'is_table': False,
                })
        return raw_blocks

    def _merge_raw_blocks(self, raw_blocks: List[Dict[str, Any]],
                          preserve_headings: bool,
                          preserve_list_items: bool) -> List[Dict[str, Any]]:
        blocks = []
        i = 0
        while i < len(raw_blocks):
            current = raw_blocks[i]
            is_atomic = (
                current['is_table']
                or (preserve_headings and current['type'] == 'heading' and current['word_count'] > 0)
                or (preserve_list_items and current['type'] == 'list_item' and current['word_count'] > 0)
            )
            if is_atomic:
                blocks.append({
                    'blocks': [current],
                    'is_table': current['is_table'],
                    'type': current['type'],
                    'text': current['text'],
                    'word_count': current['word_count'],
                })
                i += 1
                continue
            group = [current]
            merged_text = current['text']
            merged_words = current['word_count']
            j = i + 1
            while merged_words < self.min_words and j < len(raw_blocks):
                nxt = raw_blocks[j]
                if nxt['is_table']:
                    break
                if preserve_headings and nxt['type'] == 'heading' and nxt['word_count'] > 0:
                    break
                if preserve_list_items and nxt['type'] == 'list_item' and nxt['word_count'] > 0:
                    break
                group.append(nxt)
                merged_words += nxt['word_count']
                if merged_text:
                    merged_text += '\n'
                merged_text += nxt['text']
                j += 1
            first_type = group[0]['type']
            merged_type = first_type if all(b['type'] == first_type for b in group) else 'merged'
            blocks.append({
                'blocks': group,
                'is_table': False,
                'type': merged_type,
                'text': merged_text,
                'word_count': merged_words,
            })
            i = j
        return blocks

    def split(self, docx_path: str) -> List[LogicalBlock]:
        doc = Document(docx_path)
        nsmap = {'w': _W_NS}
        raw_blocks = self._build_raw_blocks(doc, nsmap)
        merged = self._merge_raw_blocks(raw_blocks, self.preserve_headings, self.preserve_list_items)
        blocks = []
        for m in merged:
            if m['is_table']:
                runs: List[RunStyle] = []
                para_style = ParagraphStyle()
            else:
                elements = [b['element'] for b in m['blocks']]
                para_objs = [p for p in doc.paragraphs if p._p in elements]
                para_objs.sort(key=lambda p: elements.index(p._p))
                runs = []
                for p in para_objs:
                    runs.extend(self._extract_run_style(run) for run in p.runs)
                para_style = self._extract_paragraph_style(para_objs[0]) if para_objs else ParagraphStyle()
            blocks.append(LogicalBlock(
                text=m['text'],
                word_count=m['word_count'],
                paragraph_style=para_style,
                runs=runs,
                block_type=m['type'],
                body_element_indices=[b['element_index'] for b in m['blocks']],
            ))
        return blocks

    def _copy_docx_with_rels(self, source_path: str, dest_path: str) -> None:
        shutil.copy2(source_path, dest_path)

    def create_block_document(self, block: LogicalBlock, output_path: str,
                               source_docx_path: str = None) -> None:
        if source_docx_path and os.path.exists(source_docx_path):
            self._copy_docx_with_rels(source_docx_path, output_path)
            doc = Document(output_path)
            body = doc.element.body
            keep_indices = set(block.body_element_indices)
            all_elements = list(body)
            if all_elements:
                last_elem = all_elements[-1]
                last_tag = last_elem.tag.split('}')[-1] if '}' in last_elem.tag else last_elem.tag
                if last_tag == 'sectPr':
                    keep_indices.add(len(all_elements) - 1)
            for idx in range(len(all_elements) - 1, -1, -1):
                if idx not in keep_indices:
                    body.remove(all_elements[idx])
            doc.save(output_path)
        else:
            new_doc = Document()
            if source_docx_path and os.path.exists(source_docx_path):
                src_doc = Document(source_docx_path)
                src_body = src_doc.element.body
                src_elements = list(src_body)
                for idx in block.body_element_indices:
                    if idx < len(src_elements):
                        new_doc.element.body.append(deepcopy(src_elements[idx]))
            else:
                pass
            new_doc.save(output_path)

    def split_to_files(self, docx_path: str, output_prefix: str) -> List[str]:
        blocks = self.split(docx_path)
        output_paths = []
        for idx, block in enumerate(blocks):
            output_path = f"{output_prefix}p{idx}d-222.docx"
            self.create_block_document(block, output_path, source_docx_path=docx_path)
            output_paths.append(output_path)
        return output_paths

    # ------------------------------------------------------------------
    # Markdown -> Word conversion helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _parse_markdown_inline(text: str) -> List[Tuple[str, Dict[str, Any]]]:
        segments: List[Tuple[str, Dict[str, Any]]] = []
        pos = 0
        for m in _MD_INLINE_RE.finditer(text):
            if m.start() > pos:
                segments.append((text[pos:m.start()], {}))
            token = m.group(0)
            if token.startswith('***') and token.endswith('***') and len(token) > 6:
                segments.append((token[3:-3], {'bold': True, 'italic': True}))
            elif token.startswith('___') and token.endswith('___') and len(token) > 6:
                segments.append((token[3:-3], {'bold': True, 'italic': True}))
            elif token.startswith('**') and token.endswith('**') and len(token) > 4:
                segments.append((token[2:-2], {'bold': True}))
            elif token.startswith('__') and token.endswith('__') and len(token) > 4:
                segments.append((token[2:-2], {'bold': True}))
            elif token.startswith('~~') and token.endswith('~~') and len(token) > 4:
                segments.append((token[2:-2], {'strike': True}))
            elif token.startswith('==') and token.endswith('==') and len(token) > 4:
                segments.append((token[2:-2], {'highlight': True}))
            elif token.startswith('++') and token.endswith('++') and len(token) > 4:
                segments.append((token[2:-2], {'underline': True}))
            elif token.startswith('~') and token.endswith('~') and len(token) > 2 and not token.startswith('~~'):
                segments.append((token[1:-1], {'subscript': True}))
            elif token.startswith('^') and token.endswith('^') and len(token) > 2:
                segments.append((token[1:-1], {'superscript': True}))
            elif token.startswith('`') and token.endswith('`') and len(token) > 2:
                segments.append((token[1:-1], {'code': True}))
            elif token.startswith('[') and token.endswith(')') and '](' in token:
                inner = token[1:-1]
                bracket_idx = inner.index('](')
                link_text = inner[:bracket_idx]
                link_url = inner[bracket_idx + 2:]
                segments.append((link_text, {'hyperlink': link_url}))
            elif token.startswith('*') and token.endswith('*') and len(token) > 2:
                segments.append((token[1:-1], {'italic': True}))
            elif token.startswith('_') and token.endswith('_') and len(token) > 2:
                segments.append((token[1:-1], {'italic': True}))
            else:
                segments.append((token, {}))
            pos = m.end()
        if pos < len(text):
            segments.append((text[pos:], {}))
        expanded: List[Tuple[str, Dict[str, Any]]] = []
        for chunk, style in segments:
            if not style:
                expanded.append((chunk, style))
                continue
            inner = DocxWorker._parse_markdown_inline(chunk)
            for inner_chunk, inner_style in inner:
                merged_style = dict(style)
                merged_style.update(inner_style)
                expanded.append((inner_chunk, merged_style))
        merged: List[Tuple[str, Dict[str, Any]]] = []
        for chunk, style in expanded:
            if merged and merged[-1][1] == style:
                merged[-1] = (merged[-1][0] + chunk, style)
            else:
                merged.append((chunk, style))
        return merged

    @staticmethod
    def _get_or_add_rPr(r_elem):
        rPr = r_elem.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            r_elem.insert(0, rPr)
        return rPr

    @classmethod
    def _set_run_bool_prop(cls, rPr, tag: str, on: bool = True) -> None:
        el = rPr.find(qn(tag))
        if on:
            if el is None:
                el = OxmlElement(tag)
                cls._insert_in_order(rPr, el, _RPR_ORDER)
            el.set(qn('w:val'), '1')
        elif el is not None:
            rPr.remove(el)

    @staticmethod
    def _insert_in_order(parent, child, order_list):
        child_tag = child.tag
        child_idx = None
        for i, tag_name in enumerate(order_list):
            if qn(tag_name) == child_tag:
                child_idx = i
                break
        if child_idx is None:
            parent.append(child)
            return
        for existing in parent:
            existing_idx = None
            for i, tag_name in enumerate(order_list):
                if qn(tag_name) == existing.tag:
                    existing_idx = i
                    break
            if existing_idx is None:
                continue
            if existing_idx > child_idx:
                existing.addprevious(child)
                return
        parent.append(child)

    def _make_run(self, text: str, style: Dict[str, Any], template_rPr) -> Any:
        r = OxmlElement('w:r')
        rPr = deepcopy(template_rPr) if template_rPr is not None else OxmlElement('w:rPr')

        # Boolean text properties
        self._set_run_bool_prop(rPr, 'w:b', style.get('bold', False))
        self._set_run_bool_prop(rPr, 'w:i', style.get('italic', False))
        self._set_run_bool_prop(rPr, 'w:strike', style.get('strike', False))

        # Underline
        if style.get('underline'):
            u = rPr.find(qn('w:u'))
            if u is None:
                u = OxmlElement('w:u')
                self._insert_in_order(rPr, u, _RPR_ORDER)
            u.set(qn('w:val'), 'single')

        # Highlight
        if style.get('highlight'):
            hl = rPr.find(qn('w:highlight'))
            if hl is None:
                hl = OxmlElement('w:highlight')
                self._insert_in_order(rPr, hl, _RPR_ORDER)
            hl.set(qn('w:val'), 'yellow')

        # Subscript
        if style.get('subscript'):
            va = rPr.find(qn('w:vertAlign'))
            if va is None:
                va = OxmlElement('w:vertAlign')
                self._insert_in_order(rPr, va, _RPR_ORDER)
            va.set(qn('w:val'), 'subscript')

        # Superscript
        if style.get('superscript'):
            va = rPr.find(qn('w:vertAlign'))
            if va is None:
                va = OxmlElement('w:vertAlign')
                self._insert_in_order(rPr, va, _RPR_ORDER)
            va.set(qn('w:val'), 'superscript')

        # FIX v7: Only set font/size if template_rPr already has them.
        # This preserves the etalon's clean style where runs don't have
        # explicit font/size unless the original template had them.
        if style.get('code'):
            font_name = _MD_CODE_FONT
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                self._insert_in_order(rPr, rFonts, _RPR_ORDER)
            for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
                rFonts.set(qn(attr), font_name)
        # Do NOT add font/size if template didn't have them — let document defaults apply

        r.append(rPr)

        if text:
            t = OxmlElement('w:t')
            t.text = text
            if text != text.strip():
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            r.append(t)
        return r

    @staticmethod
    def _strip_unwanted_pPr_elements(p_elem) -> None:
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is not None:
            for tag_name in ('w:sectPr', 'w:numPr'):
                el = pPr.find(qn(tag_name))
                if el is not None:
                    pPr.remove(el)

    def _clear_paragraph_runs(self, p_elem, nsmap) -> Any:
        """Remove all run-level content, return a template w:rPr (or None).
        FIX v7: Preserve runs that contain only w:tab/w:br (no w:t text)
        so that tab characters from the original template are kept."""
        template_rPr = None
        for r in p_elem.findall('w:r', namespaces=nsmap):
            # Check if this run has any w:t (text) element
            t_elem = r.find(qn('w:t'))
            if t_elem is None:
                # This run has no text content (e.g., w:tab, w:br) — keep it
                continue
            if template_rPr is None:
                rPr = r.find(qn('w:rPr'))
                if rPr is not None:
                    template_rPr = deepcopy(rPr)
            p_elem.remove(r)

        run_level_tags = (
            'w:hyperlink', 'w:ins', 'w:del', 'w:smartTag',
            'w:bookmarkStart', 'w:bookmarkEnd',
            'w:commentRangeStart', 'w:commentRangeEnd', 'w:commentReference',
            'w:proofErr', 'w:permStart', 'w:permEnd',
            'w:moveFrom', 'w:moveTo',
            'w:moveFromRangeStart', 'w:moveFromRangeEnd',
            'w:moveToRangeStart', 'w:moveToRangeEnd',
            'w:customXml', 'w:fldSimple', 'w:fldChar',
            'w:object', 'w:pict', 'w:txbxContent',
        )
        for tag in run_level_tags:
            for el in p_elem.findall(tag, namespaces=nsmap):
                if template_rPr is None:
                    r = el.find('.//w:r', namespaces=nsmap)
                    if r is not None:
                        rPr = r.find(qn('w:rPr'))
                        if rPr is not None:
                            template_rPr = deepcopy(rPr)
                p_elem.remove(el)
        return template_rPr

    def _apply_markdown_to_paragraph(self, p_elem, markdown_text: str, nsmap) -> None:
        """Replace paragraph content with markdown-rendered runs."""
        text = markdown_text
        text = text.replace('\r\n', '\n').replace('\r', '\n')
        text = text.replace('\\n', '\n')

        stripped = text.strip() if text else ''

        # --- Heading detection ---
        heading_level = 0
        hm = _MD_HEADING_RE.match(stripped) if stripped else None
        if hm:
            heading_level = len(hm.group(1))
            text = hm.group(2)
            stripped = text.strip()

        template_rPr = self._clear_paragraph_runs(p_elem, nsmap)

        if heading_level:
            self._apply_heading_style(p_elem, heading_level)

        # FIX v8: Do NOT apply base alignment. The template's alignment
        # is already correct and matches the etalon. Applying base_alignment
        # (JUSTIFY) would overwrite jc=right and add jc=both where there
        # should be no jc at all (e.g., "Рапорт" with jc=None).

        if not text or not text.strip():
            return

        # --- Blockquote ---
        bq = _MD_BLOCKQUOTE_RE.match(stripped)
        if bq:
            text = bq.group(1)
            self._apply_blockquote_style(p_elem)

        # --- Horizontal rule ---
        if _MD_HR_RE.match(stripped):
            self._apply_page_break(p_elem)
            return

        # --- Bullet list ---
        # FIX v9: Match against text.rstrip() (not stripped) to preserve
        # leading whitespace for nested list level detection.
        line_text = text.rstrip()
        bullet_m = _MD_BULLET_RE.match(line_text)
        if bullet_m:
            indent_str = bullet_m.group(1)
            level = len(indent_str) // 2
            text = bullet_m.group(3)
            self._apply_list_style(p_elem, bullet_type='bullet', level=level)

        # --- Numbered list ---
        num_m = _MD_NUMBERED_RE.match(line_text)
        if num_m:
            indent_str = num_m.group(1)
            level = len(indent_str) // 2
            text = num_m.group(3)
            self._apply_list_style(p_elem, bullet_type='number', level=level)

        for chunk, style in self._parse_markdown_inline(text):
            if chunk == '':
                continue
            p_elem.append(self._make_run(chunk, style, template_rPr))

    @classmethod
    def _apply_blockquote_style(cls, p_elem) -> None:
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            cls._insert_in_order(pPr, ind, _PPR_ORDER)
        ind.set(qn('w:left'), '708')
        ind.set(qn('w:hanging'), '354')
        p_rPr = pPr.find(qn('w:rPr'))
        if p_rPr is None:
            p_rPr = OxmlElement('w:rPr')
            cls._insert_in_order(pPr, p_rPr, _PPR_ORDER)
        i_elem = p_rPr.find(qn('w:i'))
        if i_elem is None:
            i_elem = OxmlElement('w:i')
            cls._insert_in_order(p_rPr, i_elem, _RPR_ORDER)
        i_elem.set(qn('w:val'), '1')

    @staticmethod
    def _apply_page_break(p_elem) -> None:
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
        pbb = pPr.find(qn('w:pageBreakBefore'))
        if pbb is None:
            pbb = OxmlElement('w:pageBreakBefore')
            pPr.insert(0, pbb)
        pbb.set(qn('w:val'), '1')

    @classmethod
    def _apply_list_style(cls, p_elem, bullet_type: str = 'bullet', level: int = 0) -> None:
        """FIX v9: Add proper w:numPr with w:ilvl and w:numId so Word renders
        actual bullet/number markers. Level 0 = top-level, level 1 = nested.
        FIX v10: Add w:jc val="both" so list items are justified (aligned
        to full width) instead of being left-aligned with a rightward shift."""
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
        # Remove existing w:ind to let numbering definition control indentation
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)
        # Remove existing w:numPr
        old_numPr = pPr.find(qn('w:numPr'))
        if old_numPr is not None:
            pPr.remove(old_numPr)
        # Add new w:numPr
        numPr = OxmlElement('w:numPr')
        cls._insert_in_order(pPr, numPr, _PPR_ORDER)
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), str(level))
        numPr.append(ilvl)
        numId = OxmlElement('w:numId')
        numId.set(qn('w:val'), '1' if bullet_type == 'bullet' else '2')
        numPr.append(numId)
        # FIX v10: Set justification to "both" (justify) for list items
        old_jc = pPr.find(qn('w:jc'))
        if old_jc is not None:
            pPr.remove(old_jc)
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'both')
        cls._insert_in_order(pPr, jc, _PPR_ORDER)

    def _ensure_numbering_definitions(self, doc) -> None:
        """FIX v9: Create numbering.xml part with bullet and numbered list
        definitions at levels 0 and 1 if it doesn't exist."""
        from docx.oxml import parse_xml
        from docx.opc.packuri import PackURI
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.parts.numbering import NumberingPart
        # Check if numbering relationship already exists
        try:
            doc.part.part_related_by(RT.NUMBERING)
            return  # Already exists
        except KeyError:
            pass
        numbering_xml = (
            f'<w:numbering xmlns:w="{_W_NS}">'
            '<w:abstractNum w:abstractNumId="0">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            '<w:lvlText w:val="\u2022"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            '<w:lvlText w:val="\u25e6"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>'
            '</w:abstractNum>'
            '<w:abstractNum w:abstractNumId="1">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
            '<w:lvlText w:val="%2)"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>'
            '</w:abstractNum>'
            '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
            '<w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>'
            '</w:numbering>'
        )
        element = parse_xml(numbering_xml)
        partname = PackURI('/word/numbering.xml')
        content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml'
        numbering_part = NumberingPart(partname, content_type, element, doc.part.package)
        doc.part.relate_to(numbering_part, RT.NUMBERING)

    @classmethod
    def _apply_heading_style(cls, p_elem, level: int) -> None:
        level = max(1, min(level, 6))
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is None:
            pStyle = OxmlElement('w:pStyle')
            cls._insert_in_order(pPr, pStyle, _PPR_ORDER)
        pStyle.set(qn('w:val'), f'Heading{level}')
        outline = pPr.find(qn('w:outlineLvl'))
        if outline is None:
            outline = OxmlElement('w:outlineLvl')
            cls._insert_in_order(pPr, outline, _PPR_ORDER)
        outline.set(qn('w:val'), str(level - 1))

    def _apply_base_alignment(self, p_elem) -> None:
        """FIX v7: Only apply base alignment if the paragraph doesn't already
        have an explicit w:jc element. This preserves original alignment from
        the template (e.g., RIGHT, inherited/None) instead of overwriting
        everything with base_alignment (JUSTIFY)."""
        a = self.base_alignment
        if a is None:
            return
        pPr = p_elem.find(qn('w:pPr'))
        # If paragraph already has explicit alignment, preserve it
        if pPr is not None:
            jc = pPr.find(qn('w:jc'))
            if jc is not None:
                return  # Keep existing alignment
        # Only set base alignment if no explicit alignment exists
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p_elem.insert(0, pPr)
        val = a.xml_value if hasattr(a, 'xml_value') else str(a)
        jc = OxmlElement('w:jc')
        self._insert_in_order(pPr, jc, _PPR_ORDER)
        jc.set(qn('w:val'), val)

    # ------------------------------------------------------------------
    # process_document
    # ------------------------------------------------------------------

    @staticmethod
    def _distribute_lines(lines: List[str], word_counts: List[int]) -> List[int]:
        P = len(word_counts)
        if P == 0:
            return []
        L = len(lines)
        if L <= 0:
            return [0] * P
        total = sum(word_counts)
        if L >= P:
            # FIX v8: Give each paragraph 1 line, all extra lines to the
            # last paragraph. This ensures clones inherit the last
            # paragraph's alignment (e.g., jc=both) instead of a random
            # paragraph's alignment (e.g., jc=right from P[0]).
            alloc = [1] * P
            rem = L - sum(alloc)
            if rem > 0:
                alloc[-1] += rem
            return alloc
        if total <= 0:
            base, rem = divmod(L, P)
            return [base + (1 if i < rem else 0) for i in range(P)]
        raw = [L * w / total for w in word_counts]
        alloc = [int(x) for x in raw]
        rem = L - sum(alloc)
        order = sorted(range(P), key=lambda i: raw[i] - alloc[i], reverse=True)
        for i in order[:rem]:
            alloc[i] += 1
        return alloc


    # ------------------------------------------------------------------
    # Markdown table support
    # ------------------------------------------------------------------

    @staticmethod
    def _is_table_start(lines: List[str], idx: int) -> bool:
        """Check if lines starting at idx form a Markdown table (header + separator)."""
        if idx + 1 >= len(lines):
            return False
        header = lines[idx].strip()
        separator = lines[idx + 1].strip()
        if '|' not in header:
            return False
        if not _MD_TABLE_SEP_RE.match(separator):
            return False
        if '|' not in separator:
            return False
        return True

    @staticmethod
    def _parse_table_row(line: str) -> List[str]:
        """Parse a Markdown table row into cell values."""
        line = line.strip()
        if line.startswith('|'):
            line = line[1:]
        if line.endswith('|'):
            line = line[:-1]
        return [cell.strip() for cell in line.split('|')]

    @staticmethod
    def _parse_table_alignment(separator: str) -> List[str]:
        """Parse column alignments from a Markdown table separator row."""
        separator = separator.strip()
        if separator.startswith('|'):
            separator = separator[1:]
        if separator.endswith('|'):
            separator = separator[:-1]
        cols = separator.split('|')
        alignments = []
        for col in cols:
            col = col.strip()
            left = col.startswith(':')
            right = col.endswith(':')
            if left and right:
                alignments.append('center')
            elif right:
                alignments.append('right')
            else:
                alignments.append('left')
        return alignments

    @staticmethod
    def _split_text_by_tables(text: str) -> List[Dict[str, Any]]:
        """Split text into segments, some of which are Markdown tables."""
        lines = text.split('\n')
        segments: List[Dict[str, Any]] = []
        i = 0
        while i < len(lines):
            if DocxWorker._is_table_start(lines, i):
                # Collect table lines
                table_lines = [lines[i], lines[i + 1]]
                alignments = DocxWorker._parse_table_alignment(lines[i + 1])
                j = i + 2
                while j < len(lines) and '|' in lines[j].strip() and lines[j].strip():
                    table_lines.append(lines[j])
                    j += 1
                segments.append({
                    'type': 'table',
                    'lines': table_lines,
                    'header': DocxWorker._parse_table_row(table_lines[0]),
                    'rows': [DocxWorker._parse_table_row(l) for l in table_lines[2:]],
                    'alignments': alignments,
                })
                i = j
            else:
                text_lines: List[str] = []
                while i < len(lines):
                    if DocxWorker._is_table_start(lines, i):
                        break
                    text_lines.append(lines[i])
                    i += 1
                segments.append({
                    'type': 'text',
                    'lines': text_lines,
                })
        return segments

    def _create_table_element(self, header: List[str], rows: List[List[str]],
                              alignments: List[str], nsmap) -> Any:
        """Create a w:tbl element from parsed Markdown table data."""
        tbl = OxmlElement('w:tbl')

        # --- Table properties ---
        tblPr = OxmlElement('w:tblPr')
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:w'), '0')
        tblW.set(qn('w:type'), 'auto')
        tblPr.append(tblW)

        # Borders
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), 'auto')
            tblBorders.append(border)
        tblPr.append(tblBorders)
        tbl.append(tblPr)

        # --- Table grid ---
        num_cols = max(len(header), max((len(r) for r in rows), default=0))
        tblGrid = OxmlElement('w:tblGrid')
        for _ in range(num_cols):
            gridCol = OxmlElement('w:gridCol')
            tblGrid.append(gridCol)
        tbl.append(tblGrid)

        def _make_cell(text: str, col_idx: int, bold: bool = False) -> Any:
            tc = OxmlElement('w:tc')
            tcPr = OxmlElement('w:tcPr')
            tc.append(tcPr)
            p = OxmlElement('w:p')
            pPr = OxmlElement('w:pPr')
            p.append(pPr)
            # Alignment
            if col_idx < len(alignments):
                align = alignments[col_idx]
                jc = OxmlElement('w:jc')
                jc.set(qn('w:val'), align)
                self._insert_in_order(pPr, jc, _PPR_ORDER)
            for chunk, style in self._parse_markdown_inline(text):
                if chunk == '':
                    continue
                run = self._make_run(chunk, style, None)
                if bold and not style.get('bold'):
                    rPr = run.find(qn('w:rPr'))
                    if rPr is not None:
                        self._set_run_bool_prop(rPr, 'w:b', True)
                p.append(run)
            tc.append(p)
            return tc

        # --- Header row ---
        tr = OxmlElement('w:tr')
        trPr = OxmlElement('w:trPr')
        tblHeader = OxmlElement('w:tblHeader')
        tblHeader.set(qn('w:val'), 'true')
        trPr.append(tblHeader)
        tr.append(trPr)
        for col_idx in range(num_cols):
            cell_text = header[col_idx] if col_idx < len(header) else ''
            tr.append(_make_cell(cell_text, col_idx, bold=True))
        tbl.append(tr)

        # --- Data rows ---
        for row_data in rows:
            tr = OxmlElement('w:tr')
            for col_idx in range(num_cols):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ''
                tr.append(_make_cell(cell_text, col_idx, bold=False))
            tbl.append(tr)

        return tbl

    def _process_block(self, block: Dict[str, Any],
                       text_processor: Callable[[str, Dict[str, Any]], str],
                       nsmap) -> None:
        elements = [b['element'] for b in block['blocks']]
        word_counts = [b['word_count'] for b in block['blocks']]

        para_info = {
            'is_heading': False,
            'is_list_item': False,
            'style_name': block['blocks'][0]['para_style'].style_name,
            'alignment': block['blocks'][0]['para_style'].alignment,
            'word_count': block['word_count'],
            'paragraph_count': len(elements),
        }
        new_text = text_processor(block['text'], para_info)
        if new_text is None:
            new_text = block['text']

        new_text = new_text.replace('\r\n', '\n').replace('\r', '\n')
        new_text = new_text.replace('\\n', '\n')

        # --- Check for Markdown tables in the new text ---
        segments = self._split_text_by_tables(new_text)
        has_tables = any(seg['type'] == 'table' for seg in segments)

        if not has_tables:
            # Original behaviour: no tables, distribute lines across paragraphs
            lines = new_text.split('\n')
            alloc = self._distribute_lines(lines, word_counts)

            pos = 0
            for elem, n in zip(elements, alloc):
                my_lines = lines[pos:pos + n]
                pos += n

                template = deepcopy(elem)
                self._strip_unwanted_pPr_elements(template)

                self._apply_markdown_to_paragraph(elem, my_lines[0] if my_lines else '', nsmap)
                anchor = elem
                for line in (my_lines[1:] if my_lines else []):
                    new_p = deepcopy(template)
                    self._strip_unwanted_pPr_elements(new_p)
                    self._apply_markdown_to_paragraph(new_p, line, nsmap)
                    anchor.addnext(new_p)
                    anchor = new_p
            return

        # --- Table-aware path: process segments in order, inserting tables ---
        elem_idx = 0
        last_anchor = None
        first_template = deepcopy(elements[0]) if elements else None
        if first_template is not None:
            self._strip_unwanted_pPr_elements(first_template)

        for seg in segments:
            if seg['type'] == 'table':
                tbl = self._create_table_element(
                    seg['header'], seg['rows'], seg['alignments'], nsmap
                )
                if last_anchor is not None:
                    last_anchor.addnext(tbl)
                    last_anchor = tbl
                elif elem_idx < len(elements):
                    elements[elem_idx].addprevious(tbl)
                    last_anchor = tbl
            else:
                for line in seg['lines']:
                    if elem_idx < len(elements):
                        elem = elements[elem_idx]
                        elem_idx += 1
                        template = deepcopy(elem)
                        self._strip_unwanted_pPr_elements(template)
                        self._apply_markdown_to_paragraph(elem, line, nsmap)
                        last_anchor = elem
                    elif first_template is not None:
                        new_p = deepcopy(first_template)
                        self._strip_unwanted_pPr_elements(new_p)
                        self._apply_markdown_to_paragraph(new_p, line, nsmap)
                        last_anchor.addnext(new_p)
                        last_anchor = new_p

    def process_document(self, input_path: str, output_path: str,
                         text_processor: Callable[[str, Dict[str, Any]], str],
                         skip_headings: bool = True,
                         skip_list_items: bool = False) -> None:
        self._copy_docx_with_rels(input_path, output_path)
        doc = Document(output_path)
        # FIX v9: Ensure numbering definitions exist before processing
        self._ensure_numbering_definitions(doc)
        nsmap = {'w': _W_NS}
        raw_blocks = self._build_raw_blocks(doc, nsmap)
        blocks = self._merge_raw_blocks(raw_blocks, skip_headings, skip_list_items)
        for block in blocks:
            if block['is_table']:
                self._process_table(block['blocks'][0]['element'], text_processor,
                                    skip_headings, skip_list_items, nsmap)
                continue
            if skip_headings and block['type'] == 'heading':
                continue
            if skip_list_items and block['type'] == 'list_item':
                continue
            self._process_block(block, text_processor, nsmap)
        self._validate_document(doc)
        doc.save(output_path)

    def _validate_document(self, doc) -> None:
        body = doc.element.body
        nsmap = {'w': _W_NS}
        bookmark_ids = set()
        for bm in body.findall('.//w:bookmarkStart', namespaces=nsmap):
            bm_id = bm.get(qn('w:id'))
            if bm_id in bookmark_ids:
                new_id = str(max([int(x) for x in bookmark_ids] + [0]) + 1)
                old_id = bm_id
                bm.set(qn('w:id'), new_id)
                bookmark_ids.add(new_id)
                for be in body.findall('.//w:bookmarkEnd', namespaces=nsmap):
                    if be.get(qn('w:id')) == old_id:
                        be.set(qn('w:id'), new_id)
                        break
            else:
                bookmark_ids.add(bm_id)
        starts = body.findall('.//w:bookmarkStart', namespaces=nsmap)
        ends = body.findall('.//w:bookmarkEnd', namespaces=nsmap)
        start_ids = {bm.get(qn('w:id')) for bm in starts}
        end_ids = {bm.get(qn('w:id')) for bm in ends}
        for bm in starts:
            if bm.get(qn('w:id')) not in end_ids:
                bm.getparent().remove(bm)
        for bm in ends:
            if bm.get(qn('w:id')) not in start_ids:
                bm.getparent().remove(bm)
        all_elements = list(body)
        if all_elements:
            last_elem = all_elements[-1]
            last_tag = last_elem.tag.split('}')[-1] if '}' in last_elem.tag else last_elem.tag
            if last_tag != 'sectPr':
                for elem in all_elements:
                    elem_tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
                    if elem_tag == 'p':
                        pPr = elem.find(qn('w:pPr'))
                        if pPr is not None:
                            sectPr = pPr.find(qn('w:sectPr'))
                            if sectPr is not None:
                                pPr.remove(sectPr)
                                body.append(sectPr)
                                break

    def _process_table(self, tbl_elem, text_processor,
                       skip_headings: bool, skip_list_items: bool, nsmap) -> None:
        # FIX v6: Process each cell independently to preserve cell boundaries.
        for row in tbl_elem.findall('.//w:tr', namespaces=nsmap):
            for cell in row.findall('.//w:tc', namespaces=nsmap):
                raw = []
                for p in cell.findall('.//w:p', namespaces=nsmap):
                    text = ''.join(t.text or '' for t in p.findall('.//w:t', namespaces=nsmap)).strip()
                    if self._is_heading_from_xml(p):
                        btype = 'heading'
                    elif self._is_list_item_from_xml(p):
                        btype = 'list_item'
                    else:
                        btype = 'paragraph'
                    raw.append({
                        'text': text,
                        'word_count': self._count_words(text),
                        'para_style': self._extract_paragraph_style_from_xml(p),
                        'runs': [],
                        'type': btype,
                        'element_index': -1,
                        'element': p,
                        'is_table': False,
                    })
                if not raw:
                    continue
                for block in self._merge_raw_blocks(raw, skip_headings, skip_list_items):
                    if skip_headings and block['type'] == 'heading':
                        continue
                    if skip_list_items and block['type'] == 'list_item':
                        continue
                    self._process_block(block, text_processor, nsmap)

    def find_context_document(self, input_path):
        doc = Document(input_path)
        nsmap = {'w': _W_NS}
        all_paragraphs = doc.element.body.findall('.//w:p', namespaces=nsmap)
        context = ''
        for para in all_paragraphs:
            context += '\n'.join(t.text or '' for t in para.findall('.//w:t', namespaces=nsmap))
        return context

    def _replace_text_in_t_elements(self, t_elements: List, new_text: str, original_lengths: List[int]) -> None:
        if not t_elements:
            return
        new_text_len = len(new_text)
        if new_text_len == 0:
            return
        num_elements = len(t_elements)
        total_original = sum(original_lengths)
        if total_original == 0 or num_elements == 1:
            t_elements[0].text = new_text
            if new_text.endswith(' ') or new_text.startswith(' '):
                t_elements[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            return
        portions = []
        for length in original_lengths:
            portion = round(new_text_len * length / total_original)
            portions.append(portion)
        diff = sum(portions) - new_text_len
        if diff > 0:
            for _ in range(diff):
                max_idx = portions.index(max(portions))
                portions[max_idx] -= 1
        elif diff < 0:
            for _ in range(-diff):
                max_idx = portions.index(max(portions))
                portions[max_idx] += 1
        current_pos = 0
        for i, t in enumerate(t_elements):
            portion_length = portions[i]
            portion_text = new_text[current_pos:current_pos + portion_length]
            current_pos += portion_length
            t.text = portion_text
            if portion_text and (portion_text.endswith(' ') or portion_text.startswith(' ')):
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        assert current_pos == new_text_len


TAG_DESCR = {
    'pos': 'Начальник бригады',
    'body': (
        'Настоящим докладыва, что я, **капитан Блохин Василий Михайлович**, '
        'в течение указанного периода выполнял следующие задачи:\n'
        '\n'
        '- Провёл инвентаризацию материально-технического обеспечения бригады\n'
        '  - проверено наличие средств индивидуальной защиты\n'
        '  - сверены данные с инвентарными карточками\n'
        '- Организовал занятия по боевой подготовке с личным составом\n'
        '  - тактическая подготовка\n'
        '  - огневая подготовка\n'
        '- Подготовил отчёт о готовности личного состава к несению службы\n'
        '\n'
        'В ходе выполнения задач были достигнуты следующие результаты:\n'
        '\n'
        '| Показатель | План | Факт | Выполнение |\n'
        '|:------------|:----:|-----:|:-----------|\n'
        '| Инвентаризация | 100% | 100% | Выполнено |\n'
        '| Боевая подготовка | 40 ч | 42 ч | Перевыполнено |\n'
        '| Готовность состава | 95% | 98% | Перевыполнено |\n'
        '\n'
        '1. Все материальные ценности учтены и внесены в реестр\n'
        '  1. проверены по наименованию\n'
        '  2. сверены по количеству\n'
        '2. Личный состав готов к выполнению задач предназначения\n'
        '3. Занятия по боевой подготовке проведены в полном объёме согласно плану'
    ),
    'date': '"15" сентября 2026 г.',
    'my_pos': 'Инженер бригады',
    'my_rank': 'капитан',
    'i_familia': 'В.Блохин'
}

def processor(text, info):
    for tag, descr in TAG_DESCR.items():
        text = text.replace('{' + '{ ' + tag + ' }' + '}', descr)
    return text

if __name__ == '__main__':
    DocxWorker().process_document(
        input_path='raport.docx',
        output_path='raport_output.docx',
        text_processor=processor
    )
