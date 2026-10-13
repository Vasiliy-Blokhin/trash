import os
from dotenv import load_dotenv
import docx
import io
from zipfile import ZipFile
import requests

from workspace.workspace_params.modules import le


load_dotenv()

# AI connection
LLM_HOST = os.getenv("LLM_HOST", '0.0.0.0')
QWEN_MODEL_SERVICE_URL = f'http://{LLM_HOST}:8835/v1'
# QWEN_MODEL_SERVICE_URL = f'http://{LLM_HOST}:8810/v1'

# AI_Manager (https://github.com/Vasiliy-Blokhin/AI_Manager): если задан
# AIM_URL, запросы к LLM идут в его OpenAI-совместимый /v1 с паролем.
AIM_URL = os.getenv("AIM_URL", '')
AIM_PASSWORD = os.getenv("AIM_PASSWORD", '')
AIM_MODEL = os.getenv("AIM_MODEL", '.')
if AIM_URL:
    QWEN_MODEL_SERVICE_URL = f'{AIM_URL.rstrip("/")}/v1'
LLM_HEADERS = {'X-API-Password': AIM_PASSWORD} if AIM_PASSWORD else None
# AI_Manager ставит запросы в очередь, поэтому таймаут больше
LLM_TIMEOUT = float(os.getenv("LLM_TIMEOUT", '900'))
LLM_RETRIES = int(os.getenv("LLM_RETRIES", '1'))

QWEN_SETTINGS = lambda message_list = None : {
    'full_url': f'{QWEN_MODEL_SERVICE_URL}/chat/completions',
    'short_url': QWEN_MODEL_SERVICE_URL,
    'request_data': {
        "model": AIM_MODEL if AIM_URL else ".",
        "messages": message_list,
        "temperature": 0.1,
        "top_p": 0.9,
    }
}

EMBEDDING_SERVICE_URL = f'http://{LLM_HOST}:9000/v1'
EMBEDDING_MODEL_URL = f'{EMBEDDING_SERVICE_URL}/embeddings'
EMBEDDING_SETTINGS = lambda text = None : {
        'full_url': EMBEDDING_MODEL_URL,
        'short_url': EMBEDDING_SERVICE_URL,
        'request_data': {
                "model": ".",
                "input": text,
            }
    }

TOKENIZER_MODEL_URL = f'http://{LLM_HOST}:8810/tokenize'
TOKENIZER_SETTINGS = lambda text: {
        'full_url': TOKENIZER_MODEL_URL,
        'request_data': {
                "model": ".",
                "input": text,
            }
    }

OCR_SERVICE_URL = lambda key: f'http://{LLM_HOST}:9999/ocr/{key}'

GIGACHAT_CERT = './workspace/ai_connection/cert/gigachat.crt'
GIGACHAT_KEY = './workspace/ai_connection/cert/gigachat.key'
# -----------------------------------------------------------------------------

# DB connetction
DB_CONFIG = {
    'host': os.getenv("DATABASES_HOST", '0.0.0.0'),
    'port': os.getenv("DATABASES_PORT", '5432'),
    'database': os.getenv("DATABASES_NAME", 'rel_db'),
    'user': os.getenv("DATABASES_USER", 'postgres'),
    'password': os.getenv("DATABASES_PASSWORD", '12345678'),
    'min_connections': 1,
    'max_connections': 20,
    'statement_timeout': "30s"
}
# -----------------------------------------------------------------------------

# TEST request data from metodic
TEST_FILE_PATH = './in_data/input.docx' # можно 
# TEST_FILE_PATH = './in_data/Власенкова_43_8.docx' # нельзя


def _generate_test_file_text_for_request(docx_info):
    """Имитация получения информации с файла."""
    try:
        file_text = ""
        document = docx.Document(docx_info)
        for paragraph in document.paragraphs:
            file_text += paragraph.text + " "

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph_in_cell in cell.paragraphs:
                        file_text += paragraph_in_cell.text + " "

        return file_text.strip()

    except Exception as e:
        le(f"Ошибка при формировании текста - {e}")
        return str(e) 


REQUEST_PARAMS = {
    'request_data': {
        'model': '.',
        'messages': [
            {
                'role': 'user',
                'content': 'Запрос с окна чата'
            }
        ],
        'temperature': 0.1,
        'top_p': 0.9
    },
    'full_url': 'http://0.0.0.0:8835/v1/chat/completions',
    'short_url': 'http://0.0.0.0:8835/v1',
    'file_content': _generate_test_file_text_for_request(TEST_FILE_PATH),
    # в file_content передается ссылка на документ, который можно получить 
    # через get запрос, ниже представлен базовый метод получения информации из
    # документа docx, полученного таким путем
    'chat_id': 'b401e3c7b1af40aeab24065c8a556a8d',
    'task_id': 'b1c9119539ad-4ba1-9f3e-f7bca87b2054', 
    'state_agent': None,
    'group_key': '28',
    'search_in_db': False,
    'user': {
        'id': '66794dfdfd3e4f928daa7d275299ca5b',
        'username': 'blokhin_vm',
        'id_db': None,
        'user_db': 'blokhin_vm_1ae8edc2425744a29cebc3e86e4c02ac',
        'person': None,
        'usertype': 'Администратор',
        'id_objuid': 1,
        'full_name':'blokhin_vm',
        'is_active': True,
        'is_admin': True,
        'workgroup_code': '28',
        'workgroup_pk': '58b3d88f8a3141139aa1d9e9d12dff65',
        'sessionid': None,
        'liberty_prime_sessionid': 'vyaj07pzdasps4l7wpf2wqy49hhzal7k'
    }
}

def _get_file_from_url(url):
    """
    Ознакомительный метод по получению документа из запроса и преобразование 
    в текстовый вид (docx)
    """
    response = requests.get(url=url)

    try:
        file_like_object = io.BytesIO(response.content)

        with ZipFile(file_like_object) as myzip:
            if "word/document.xml" not in myzip.namelist():
                raise Exception("Это невалидный DOCX файл.")

            # либо file_like_object можешь сохранить файлом, или 
            # сохранить бинарник значением класса
            return _generate_test_file_text_for_request(file_like_object)

    except Exception as e:
        le(f"Ошибка обработки файла: {e}")
# -----------------------------------------------------------------------------

# other params
URL_BACKEND = 'pass'
STATE = 'pass'